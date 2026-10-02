"""Build an app-competition dataset from NPCI's monthly UPI app workbooks.

Usage: python build.py --source path/to/monthly/xlsx/files
"""

from __future__ import annotations

import argparse
import csv
import re
from calendar import month_abbr
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "data" / "processed"
FILE_DATE = re.compile(r"(20\d{2})-([A-Za-z]{3})\.xlsx$")
MONTH_NUMBER = {name.lower(): month for month, name in enumerate(month_abbr) if name}
LEADERS = ("PhonePe", "Google Pay", "Paytm")


def parse_number(value: object) -> float | None:
    """Keep a visibly unreadable source cell missing instead of making it zero."""
    if value is None or str(value).strip() == "" or set(str(value).strip()) == {"#"}:
        return None
    if str(value).strip() in {"-", "–", "—"}:
        return 0.0
    return float(str(value).replace(",", "").strip())


def app_name(value: object) -> str:
    name = re.sub(r"\s*#\s*$", "", str(value).strip())
    names = {
        "Phone Pe": "PhonePe", "Phonepe": "PhonePe",
        "Paytm Payments Bank App": "Paytm", "Paytm (OCL)": "Paytm",
        "Paytm (OCL )": "Paytm", "Cred": "CRED",
        "WhatsApp": "WhatsApp Pay", "Whatsapp": "WhatsApp Pay",
    }
    return names.get(name, name)


def read_workbook(path: Path) -> list[dict]:
    match = FILE_DATE.search(path.name)
    if match is None:
        raise ValueError(f"Workbook filename has no year/month: {path.name}")
    year = int(match.group(1))
    month = MONTH_NUMBER[match.group(2).lower()]
    date = f"{year:04d}-{month:02d}-01"
    workbook = load_workbook(path, read_only=True, data_only=True)
    header_index = None
    total_column = None
    observations = []
    try:
        for excel_row, row in enumerate(workbook.active.iter_rows(values_only=True), start=1):
            if header_index is None:
                if "Application Name" in row:
                    labels = [str(cell).strip() if cell is not None else "" for cell in row]
                    total_column = labels.index("Total")
                    if total_column not in (8, 10):
                        raise ValueError(f"Unexpected table shape in {path.name}")
                    header_index = excel_row
                continue
            serial = str(row[0]).strip() if row and row[0] is not None else ""
            if not re.fullmatch(r"\d+(?:\.0)?", serial):
                continue
            if len(row) <= total_column + 1 or row[1] is None:
                raise ValueError(f"Incomplete app record at {path.name}:{excel_row}")
            volume = parse_number(row[total_column])
            value = parse_number(row[total_column + 1])
            if volume is None:
                raise ValueError(f"Missing transaction volume at {path.name}:{excel_row}")
            observations.append({
                "date": date, "year": year, "month": month,
                "app": app_name(row[1]), "volume_mn": volume,
                "value_cr": value, "source_file": path.name,
                "source_row": excel_row,
            })
    finally:
        workbook.close()
    if not observations:
        raise ValueError(f"No data rows in {path.name}")
    return observations


def write_csv(name: str, fields: list[str], records: list[dict]) -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    with (OUTPUT / name).open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def build(source: Path) -> None:
    files = sorted(source.glob("*.xlsx"))
    if not files:
        raise FileNotFoundError(f"No .xlsx workbooks found in {source}")
    raw = [observation for path in files for observation in read_workbook(path)]
    dates = {row["date"] for row in raw}
    if len(dates) != len(files):
        raise ValueError("Expected one source workbook per month")

    # A few workbooks list the same app more than once. Combine those rows.
    app_totals: dict[tuple[str, str], dict] = {}
    for row in raw:
        key = (row["date"], row["app"])
        if key not in app_totals:
            app_totals[key] = {
                "date": row["date"], "year": row["year"], "month": row["month"],
                "app": row["app"], "volume_mn": 0.0, "value_cr": 0.0,
                "value_complete": 1,
            }
        item = app_totals[key]
        item["volume_mn"] += row["volume_mn"]
        if row["value_cr"] is None:
            item["value_complete"] = 0
        else:
            item["value_cr"] += row["value_cr"]
    app_month = []
    for _, item in sorted(app_totals.items()):
        item["volume_mn"] = round(item["volume_mn"], 2)
        item["value_cr"] = round(item["value_cr"], 2) if item["value_complete"] else None
        app_month.append(item)

    by_date: dict[str, list[dict]] = defaultdict(list)
    for item in app_month:
        by_date[item["date"]].append(item)
    month_summary = []
    leaders_monthly = []
    for date, apps in sorted(by_date.items()):
        total = sum(row["volume_mn"] for row in apps)
        if total <= 0:
            raise ValueError(f"Zero market volume for {date}")
        names = {row["app"]: row["volume_mn"] for row in apps}
        named = sorted((row["volume_mn"] for row in apps if row["app"].lower() not in {"other", "others"}), reverse=True)
        shares = {name: names.get(name, 0.0) / total for name in LEADERS}
        month_summary.append({
            "date": date, "year": int(date[:4]), "month": int(date[5:7]),
            "reported_volume_mn": round(total, 2),
            "listed_apps": sum(row["app"].lower() not in {"other", "others"} for row in apps),
            "top3_share": round(sum(named[:3]) / total, 6),
            "phonepe_share": round(shares["PhonePe"], 6),
            "googlepay_share": round(shares["Google Pay"], 6),
            "paytm_share": round(shares["Paytm"], 6),
            "missing_value_apps": sum(not row["value_complete"] for row in apps),
        })
        leader_volume = 0.0
        for name in LEADERS:
            volume = names.get(name, 0.0)
            leader_volume += volume
            leaders_monthly.append({"date": date, "year": int(date[:4]), "app_group": name,
                                    "volume_mn": round(volume, 2), "share": round(volume / total, 6)})
        other = total - leader_volume
        leaders_monthly.append({"date": date, "year": int(date[:4]), "app_group": "All other apps",
                                "volume_mn": round(other, 2), "share": round(other / total, 6)})

    yearly_totals: dict[tuple[int, str], float] = defaultdict(float)
    yearly_market: dict[int, float] = defaultdict(float)
    for row in leaders_monthly:
        yearly_totals[(row["year"], row["app_group"])] += row["volume_mn"]
        yearly_market[row["year"]] += row["volume_mn"]
    yearly = [
        {"year": year, "app_group": group, "volume_mn": round(volume, 2),
         "share": round(volume / yearly_market[year], 6)}
        for (year, group), volume in sorted(yearly_totals.items())
    ]
    by_calendar_month: dict[int, list[float]] = defaultdict(list)
    for row in month_summary:
        by_calendar_month[row["month"]].append(row["reported_volume_mn"])
    seasonality = [
        {"month": month, "month_name": month_abbr[month],
         "mean_reported_volume_mn": round(sum(values) / len(values), 2),
         "years_observed": len(values)}
        for month, values in sorted(by_calendar_month.items())
    ]
    latest = max(dates)
    latest_total = next(row["reported_volume_mn"] for row in month_summary if row["date"] == latest)
    ranking = [
        {"rank": rank, "app": row["app"], "volume_mn": row["volume_mn"],
         "share": round(row["volume_mn"] / latest_total, 6), "date": latest}
        for rank, row in enumerate(sorted(by_date[latest], key=lambda item: item["volume_mn"], reverse=True)[:20], start=1)
    ]

    write_csv("app_month.csv", ["date", "year", "month", "app", "volume_mn", "value_cr", "value_complete"], app_month)
    write_csv("market_month.csv", list(month_summary[0]), month_summary)
    write_csv("leaders_month.csv", list(leaders_monthly[0]), leaders_monthly)
    write_csv("leaders_year.csv", list(yearly[0]), yearly)
    write_csv("seasonality.csv", list(seasonality[0]), seasonality)
    write_csv("latest_top20.csv", list(ranking[0]), ranking)
    print(f"Built {len(app_month)} app-month records from {len(files)} NPCI workbooks")
    print(f"Latest month: {latest}; reported app volume: {latest_total:,.2f} million")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "data" / "raw",
                        help="Folder containing NPCI monthly app .xlsx files")
    build(parser.parse_args().source)
